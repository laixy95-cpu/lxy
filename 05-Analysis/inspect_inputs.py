from pathlib import Path
import json
import re

import pandas as pd
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "source_snapshot"
AUDIT = ROOT / "data_audit"
AUDIT.mkdir(parents=True, exist_ok=True)


def docx_text(path: Path) -> str:
    doc = Document(path)
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text.replace("\n", " | ") for cell in row.cells))
    return "\n".join(parts)


inventory = []
for path in sorted(SRC.iterdir()):
    item = {"file": path.name, "bytes": path.stat().st_size, "suffix": path.suffix.lower()}
    if path.suffix.lower() == ".xlsx":
        book = pd.ExcelFile(path)
        item["sheets"] = []
        for sheet in book.sheet_names:
            df = pd.read_excel(path, sheet_name=sheet)
            item["sheets"].append({
                "name": sheet,
                "rows": int(df.shape[0]),
                "cols": int(df.shape[1]),
                "columns": [str(c) for c in df.columns],
            })
    elif path.suffix.lower() == ".docx":
        text = docx_text(path)
        (AUDIT / f"{path.stem}_extracted.txt").write_text(text, encoding="utf-8")
        item["characters_extracted"] = len(text)
        item["keyword_hits"] = {
            key: len(re.findall(key, text, flags=re.I))
            for key in ["ETRI", "TDLoss", "AccessElec", "Inflation", "maintenance", "policy"]
        }
    inventory.append(item)

(AUDIT / "input_inventory.json").write_text(
    json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(json.dumps(inventory, ensure_ascii=False, indent=2))
